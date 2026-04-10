import openpyxl
import pandas as pd
import os
import sys
from calendar import monthrange

# 1. Configuración de Ingesta y Rutas
ANIO = "2025"
MESES = [f"{i:02d}" for i in range(1, 13)] # Generar 01, 02, ..., 12
BASE_DIR = r"C:\DESARROLLO\EXTRACCION_DATOS_REM\data_rem"
RUTA_DICCIONARIO = r"C:\DESARROLLO\EXTRACCION_DATOS_REM\diccionario_centros\COD_CENTROS_SALUD.CSV"
RUTA_SALIDA_CARPETA = r"C:\DESARROLLO\EXTRACCION_DATOS_REM\resumen\BM"

# Configuración de celda Serie BM (Curación Avanzada Pie Diabético)
HOJA_FUENTE = "BM18A"

# 2. Carga de Diccionario de Centros
try:
    df_centros = pd.read_csv(RUTA_DICCIONARIO)
except Exception as e:
    print(f"Error al cargar el diccionario de centros: {e}")
    sys.exit(1)

datos_maestros = []

print(f"--- Iniciando Extracción Masiva Serie BM - Mensual ({ANIO}) ---")

# 3. Validación de Existencia y Procesamiento Persistente
for mes in MESES:
    print(f"\nVerificando Periodo: {mes}/{ANIO}...")
    
    # Calcular el último día del mes para la fecha de corte
    ultimo_dia = monthrange(int(ANIO), int(mes))[1]
    fecha_corte = f"{ultimo_dia}/{mes}/{ANIO}"
    
    for _, centro in df_centros.iterrows():
        codigo = str(centro['COD_CENTRO'])
        nombre = centro['NOMBRE']
        
        # Construir ruta esperada del archivo BM
        archivo_nombre = f"{codigo}BM.xlsm"
        ruta_archivo = os.path.join(BASE_DIR, ANIO, "BM", mes, archivo_nombre)
        
        # Validar existencia del archivo
        if not os.path.exists(ruta_archivo):
            # Reintento con extensión .xls
            ruta_archivo_alt = ruta_archivo.replace(".xlsm", ".xls")
            if os.path.exists(ruta_archivo_alt):
                ruta_archivo = ruta_archivo_alt
            else:
                print(f"\n[ERROR DE INGESTA] Proceso detenido.")
                print(f"Falta archivo Serie BM: {archivo_nombre}")
                print(f"Centro: {codigo} - {nombre}")
                print(f"Periodo faltante: {mes}/{ANIO}")
                sys.exit(1)
        
        # Extracción inteligente del dato (Búsqueda por palabra clave)
        print(f" > Extrayendo: {codigo} - {nombre}...", end="\r")
        try:
            wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
            
            # Verificar si la hoja existe
            if HOJA_FUENTE not in wb.sheetnames:
                print(f"\n[ERROR] Hoja {HOJA_FUENTE} no encontrada en {archivo_nombre}")
                continue
                
            hoja = wb[HOJA_FUENTE]
            
            # Búsqueda inteligente en columna B
            item_encontrado = False
            
            # Recorremos la columna B para encontrar la fila del item
            for row in hoja.iter_rows(min_col=2, max_col=2): # Columna B
                celda_b = row[0]
                if celda_b.value and "Curación Avanzada Pie Diabético" in str(celda_b.value):
                    fila_idx = celda_b.row
                    item_encontrado = True
                    
                    # Definición de columnas de extracción (D, E, F) y sus nombres de centro
                    mapeo_centros = [
                        {"col": 4, "nombre": "SAPU/SAR/SUR"},
                        {"col": 5, "nombre": "Resto Establecimientos APS"},
                        {"col": 6, "nombre": "Compra de Servicio"}
                    ]
                    
                    for item_c in mapeo_centros:
                         val_celda = hoja.cell(row=fila_idx, column=item_c["col"]).value
                         valor_final = float(val_celda) if isinstance(val_celda, (int, float)) else 0
                         
                         # Construcción del registro
                         datos_maestros.append({
                             "Fecha Corte": fecha_corte,
                             "Codigo DEIS": codigo,
                             "Nombre Centro": nombre,
                             "Serie": "BM",
                             "Hoja": HOJA_FUENTE,
                             "Item": "Curación Avanzada Pie Diabético",
                             "Centro": item_c["nombre"],
                             "Cantidad": valor_final,
                             "Ruta Origen": ruta_archivo
                         })
                    break
            
            if not item_encontrado:
                 print(f"\n[ADVERTENCIA] No se encontró la etiqueta en: {archivo_nombre}")
            
        except Exception as e:
            print(f"\nError al procesar el archivo {archivo_nombre}: {e}")

print(f"\n\n--- Extracción Serie BM finalizada con éxito ---")

# 4. Consolidación y Exportación
if datos_maestros:
    df_final = pd.DataFrame(datos_maestros)
    
    if not os.path.exists(RUTA_SALIDA_CARPETA):
        os.makedirs(RUTA_SALIDA_CARPETA)
    
    ruta_xlsx = os.path.join(RUTA_SALIDA_CARPETA, "bm_curacion_avanzada_pie_diabetico.xlsx")
    df_final.to_excel(ruta_xlsx, index=False, sheet_name="MISCELANEOS")
    
    print(f"Producto consolidado generado en: {ruta_xlsx}")
    print(f"Total de registros procesados: {len(df_final)}")
else:
    print("No se encontraron datos para procesar.")